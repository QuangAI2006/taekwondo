from __future__ import annotations

import csv
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = Path(r"<DESKTOP>/Sensors_Submission_Package_20260524_081544")
SOURCE_DOCX = PACKAGE / "Sensors_revised_manuscript.docx"
OUT_DOCX = PACKAGE / "Sensors_revised_manuscript_experiment_checked.docx"

OUT = ROOT / "outputs"
REV = OUT / "manuscript_revisions"
TABLES = REV / "tables"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def read_md_body(path: Path) -> str:
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    if lines and lines[0].startswith("#"):
        lines = lines[1:]
    return "\n".join(lines).strip()


def style(doc: Document, name: str, fallback: str = "Normal") -> str:
    names = {s.name for s in doc.styles}
    return name if name in names else fallback


def set_para_text(doc: Document, starts: str, text: str) -> None:
    for p in doc.paragraphs:
        if p.text.strip().startswith(starts):
            p.text = text
            return
    raise ValueError(f"Paragraph not found: {starts}")


def find_para(doc: Document, starts: str):
    for p in doc.paragraphs:
        if p.text.strip().startswith(starts):
            return p
    raise ValueError(f"Paragraph not found: {starts}")


def body_blocks(doc: Document):
    body = doc.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield child, "".join(child.itertext()).strip()
        elif isinstance(child, CT_Tbl):
            yield child, ""


def delete_section(doc: Document, start_text: str, end_text: str) -> None:
    blocks = list(body_blocks(doc))
    start = None
    end = None
    for i, (_el, text) in enumerate(blocks):
        if start is None and text.startswith(start_text):
            start = i
        if start is not None and text.startswith(end_text):
            end = i
            break
    if start is None or end is None:
        raise ValueError(f"Could not find section range: {start_text} -> {end_text}")
    body = doc.element.body
    for el, _text in blocks[start:end]:
        body.remove(el)


def insert_paragraph(anchor, text: str = "", style_name: str | None = None):
    p = anchor.insert_paragraph_before(text)
    if style_name:
        p.style = style_name
    return p


def insert_heading(doc: Document, anchor, text: str, level: int) -> None:
    if level == 1:
        st = style(doc, "MDPI_2.1_heading1")
    elif level == 2:
        st = style(doc, "MDPI_2.2_heading2")
    else:
        st = style(doc, "MDPI_2.3_heading3")
    insert_paragraph(anchor, text, st)


def insert_body_paragraph(doc: Document, anchor, text: str) -> None:
    if not text.strip():
        return
    insert_paragraph(anchor, text, style(doc, "MDPI_3.1_text"))


def insert_caption(doc: Document, anchor, text: str, figure: bool = False) -> None:
    st = style(doc, "MDPI_5.1_figure_caption" if figure else "MDPI_4.1_table_caption")
    insert_paragraph(anchor, text, st)


def insert_page_break(anchor) -> None:
    p = anchor.insert_paragraph_before("")
    p.add_run().add_break(WD_BREAK.PAGE)


def insert_table(doc: Document, anchor, rows: list[dict[str, Any]], fields: list[str], headers: list[str]) -> None:
    table = doc.add_table(rows=1, cols=len(fields))
    table.style = style(doc, "MDPI_4.1_three_line_table", "Table Grid")
    table.autofit = True
    for j, header in enumerate(headers):
        table.cell(0, j).text = header
    for row in rows:
        cells = table.add_row().cells
        for j, field in enumerate(fields):
            cells[j].text = str(row.get(field, ""))
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.style = style(doc, "MDPI_4.2_table_body", "Normal")
    anchor._p.addprevious(table._tbl)


def insert_picture(doc: Document, anchor, image_path: Path, width_inches: float) -> None:
    p = doc.add_paragraph()
    p.style = style(doc, "MDPI_5.2_figure", "Normal")
    p.alignment = 1
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    anchor._p.addprevious(p._p)


def insert_paragraph_after(doc: Document, anchor, text: str, style_name: str | None = None):
    p = doc.add_paragraph(text)
    if style_name:
        p.style = style_name
    anchor._p.addnext(p._p)
    return p


def insert_copied_table(anchor, table_xml) -> None:
    anchor._p.addprevious(deepcopy(table_xml))


def paragraphs_from_markdown(text: str) -> list[tuple[str, int | None]]:
    out: list[tuple[str, int | None]] = []
    buf: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("## "):
            if buf:
                out.append((" ".join(buf), None))
                buf = []
            out.append((stripped[3:], 2))
        elif stripped.startswith("# "):
            if buf:
                out.append((" ".join(buf), None))
                buf = []
            out.append((stripped[2:], 1))
        elif not stripped:
            if buf:
                out.append((" ".join(buf), None))
                buf = []
        else:
            buf.append(stripped)
    if buf:
        out.append((" ".join(buf), None))
    return out


def add_section4(doc: Document, anchor, fig3_table_xml, fig3_caption: str) -> None:
    table1 = read_csv(TABLES / "table1_acquisition_protocol.csv")
    table2 = read_csv(TABLES / "table2_dataset_split_summary.csv")
    table3 = read_csv(TABLES / "table3_model_training_configuration.csv")
    table4 = read_csv(TABLES / "table4_backbone_comparison.csv")
    table5 = read_csv(TABLES / "table5_temporal_ablation.csv")
    table6 = read_csv(TABLES / "table6_classwise_final_model.csv")
    table7 = read_csv(TABLES / "table7_runtime.csv")
    table8 = read_csv(TABLES / "table8_human_ai_agreement.csv")

    insert_heading(doc, anchor, "4. Experiments and Analysis", 1)
    insert_heading(doc, anchor, "4.1. Dataset and Current Validation Protocol", 2)
    insert_body_paragraph(
        doc,
        anchor,
        "The TKD-Kick3 acquisition protocol included 140 first-year students from four teaching classes. Each student was asked to record two videos for each of three fundamental kicks, giving 840 planned smartphone videos. After cleaning, the active repository contained 765 retained three-class pose sequences. Side-kick videos were excluded from the present three-class study.",
    )
    insert_body_paragraph(
        doc,
        anchor,
        "A critical audit found that participant identifiers were not available in the active sequence files or in a separate metadata table. Therefore, a subject-independent split could not be generated. The current experiments use the existing cleaned file-level training and validation split and should be interpreted as sample-level validation rather than participant-independent testing.",
    )
    insert_caption(doc, anchor, "Table 1. Smartphone-video acquisition protocol.")
    insert_table(doc, anchor, table1, ["item", "value"], ["Item", "Description"])
    insert_caption(doc, anchor, "Table 2. Dataset inclusion, exclusion, and current split summary.")
    insert_table(doc, anchor, table2, ["split_or_stage", "participants", "front", "roundhouse", "axe", "total", "note"], ["Stage/Split", "Participants", "Front", "Roundhouse", "Axe", "Total", "Note"])

    insert_heading(doc, anchor, "4.2. Model Configurations and Training", 2)
    insert_body_paragraph(
        doc,
        anchor,
        "All recognition models used 96-frame sequences with 52-dimensional frame features. Training was repeated with seeds 0, 1, and 2. Adam was used with an initial learning rate of 1e-3, batch size 8, dropout 0.2, and 40 epochs. The Transformer used label smoothing of 0.05; the LSTM and GCN baseline scripts did not apply label smoothing.",
    )
    insert_caption(doc, anchor, "Table 3. Model architecture and training configuration.")
    insert_table(doc, anchor, table3, ["model", "input", "main_dimensions", "dropout", "parameters", "training"], ["Model", "Input", "Architecture", "Dropout", "Parameters", "Training"])

    insert_heading(doc, anchor, "4.3. Recognition Results", 2)
    insert_body_paragraph(
        doc,
        anchor,
        "The strongest recognition result was obtained by the BiLSTM with uniform temporal resampling and argmax inference. This configuration achieved 75.3 +/- 2.3% accuracy, 74.3 +/- 2.3% macro-F1, and 74.5 +/- 2.3% balanced accuracy across three seeds. LSTM-PA-Sweep was also run and reached 71.2 +/- 1.0% accuracy, so phase-aligned sweep inference did not improve the LSTM under the current split.",
    )
    insert_body_paragraph(
        doc,
        anchor,
        "The Transformer with combined phase-aligned resampling and sweep inference reached 67.3 +/- 1.7% accuracy, compared with 62.2 +/- 2.4% for Transformer-UNI-Argmax. The present results therefore support LSTM-UNI as the selected recognition backbone rather than the Transformer.",
    )
    insert_caption(doc, anchor, "Table 4. Recognition backbone comparison under the current sample-level validation protocol.")
    insert_table(doc, anchor, table4, ["model", "setting", "accuracy", "macro_f1", "weighted_f1", "balanced_accuracy", "n_runs"], ["Model", "Configuration", "Accuracy (%)", "Macro-F1 (%)", "Weighted-F1 (%)", "Balanced Acc. (%)", "Seeds"])
    insert_caption(doc, anchor, "Table 5. Ablation of temporal normalization and inference strategy for the Transformer model.")
    insert_table(doc, anchor, table5, ["configuration", "deployable", "accuracy", "macro_f1", "balanced_accuracy", "interpretation"], ["Configuration", "Deployable", "Accuracy (%)", "Macro-F1 (%)", "Balanced Acc. (%)", "Interpretation"])
    insert_caption(doc, anchor, "Table 6. Class-wise performance of the selected LSTM-UNI configuration.")
    insert_table(doc, anchor, table6, ["class", "precision", "recall", "f1", "support"], ["Class", "Precision (%)", "Recall (%)", "F1 (%)", "Support"])
    if fig3_table_xml is not None:
        insert_page_break(anchor)
        insert_copied_table(anchor, fig3_table_xml)
        insert_caption(doc, anchor, fig3_caption, figure=True)
    insert_picture(doc, anchor, OUT / "figures" / "confusion_matrix_final_model.png", 5.4)
    insert_caption(doc, anchor, "Figure 4. Row-normalized confusion matrix of the selected LSTM-UNI recognition model.", figure=True)

    insert_heading(doc, anchor, "4.4. Runtime Performance", 2)
    insert_body_paragraph(
        doc,
        anchor,
        "Runtime was measured on the current desktop CPU prototype using 30 validation videos, with 10 videos per class. The total processing time excluding feedback rendering was 7.53 +/- 4.53 s per video. Recognition inference itself required only 0.0030 +/- 0.0014 s per video; pose extraction was the dominant computational stage.",
    )
    insert_caption(doc, anchor, "Table 7. Runtime performance on a desktop CPU prototype.")
    insert_table(doc, anchor, table7, ["stage", "n", "mean_seconds", "std_seconds", "min_seconds", "max_seconds"], ["Stage", "N", "Mean (s/video)", "SD", "Min", "Max"])

    insert_heading(doc, anchor, "4.5. Explainable Scoring and Human-AI Agreement", 2)
    insert_body_paragraph(
        doc,
        anchor,
        "The scoring module produced a 10-point score consisting of 4 points for accuracy and 6 points for expressiveness. The expression component used speed, height, smoothness, and a pose-derived kinematic vigor index. This index should not be interpreted as biomechanical power because no force, torque, or three-dimensional dynamics were measured.",
    )
    insert_caption(doc, anchor, "Table 8. Human-AI agreement on the expert-annotated subset.")
    insert_table(doc, anchor, table8, ["metric", "result"], ["Metric", "Result"])
    insert_page_break(anchor)
    insert_picture(doc, anchor, OUT / "expert_agreement" / "bland_altman_plot.png", 5.2)
    insert_caption(doc, anchor, "Figure 5. Bland-Altman plot comparing system scores with mean expert scores.", figure=True)
    insert_body_paragraph(
        doc,
        anchor,
        "These results provide preliminary evidence that the rule-based scoring layer is directionally consistent with expert evaluation, but the expert subset is small and cannot currently be linked to a confirmed subject-independent test set.",
    )


def add_markdown_section(doc: Document, anchor, markdown_path: Path, skip_title: bool = False) -> None:
    body = read_md_body(markdown_path)
    for text, level in paragraphs_from_markdown(body):
        if skip_title and level == 1:
            continue
        if level is not None:
            insert_heading(doc, anchor, text, level)
        else:
            insert_body_paragraph(doc, anchor, text)


def main() -> None:
    doc = Document(SOURCE_DOCX)

    # Preserve existing visual feedback example table before replacing Section 4.
    fig3_table_xml = deepcopy(doc.tables[3]._tbl) if len(doc.tables) > 3 else None
    fig3_caption = "Figure 3. Visual feedback examples generated by the proposed system in the human-AI agreement study."
    for p in doc.paragraphs:
        if p.text.strip().startswith("Figure 3. Visual feedback examples"):
            fig3_caption = p.text.strip()
            break

    abstract = read_md_body(REV / "abstract_updated.md")
    abstract = abstract.replace("# Updated Abstract", "").strip()
    set_para_text(doc, "With the rapid proliferation", abstract)
    set_para_text(doc, "Keywords:", "Keywords: smartphone video; pose estimation; action recognition; action quality assessment; taekwondo; explainable feedback")

    set_para_text(
        doc,
        "To address subjective assessment",
        "To address subjective assessment, delayed feedback, and the difficulty of producing quantified, traceable evaluations in taekwondo fundamentals training, we propose a smartphone-video-based assessment pipeline integrating pose extraction, kick recognition, rule-based scoring, and visual feedback. The current implementation is a desktop-CPU prototype that analyzes smartphone-captured videos rather than a verified on-device mobile deployment.",
    )
    set_para_text(
        doc,
        "For recognition, we develop",
        "For recognition, we evaluate three temporal backbones: a bidirectional LSTM [20], a GCN-style baseline [9,21], and a lightweight Transformer encoder [5]. The final backbone is selected from the reproduced validation results rather than fixed a priori. In the current experiment matrix, LSTM with uniform temporal resampling and argmax inference provides the strongest recognition performance.",
    )
    set_para_text(
        doc,
        "We propose a mobile-oriented framework",
        "We propose a smartphone-video-based framework for recognizing and quantitatively evaluating fundamental taekwondo kicks. As illustrated in Figure 1, the system integrates four core modules: skeletal data acquisition and feature engineering, temporal recognition, rule-based scoring, and visual feedback generation. The current prototype runs on a desktop CPU while using smartphone-captured videos as input.",
    )
    set_para_text(doc, "Transformer-based temporal modeling", "Recognition backbones and temporal modeling")

    delete_section(doc, "4. Experiments and Analysis", "5. Discussion")
    anchor5 = find_para(doc, "5. Discussion")
    add_section4(doc, anchor5, fig3_table_xml, fig3_caption)

    delete_section(doc, "5. Discussion", "6. Conclusion")
    anchor6 = find_para(doc, "6. Conclusion")
    insert_heading(doc, anchor6, "5. Discussion", 1)
    insert_heading(doc, anchor6, "5.1. Recognition Backbone and Temporal Alignment", 2)
    insert_body_paragraph(
        doc,
        anchor6,
        "The reproduced experiment matrix changes the interpretation of the recognition module. Although the Transformer benefits from the combined phase-aligned resampling and sweep-inference configuration, the strongest current result is obtained by LSTM-UNI. The manuscript should therefore present the Transformer as an evaluated comparison model rather than as the final system backbone.",
    )
    insert_body_paragraph(
        doc,
        anchor6,
        "The temporal ablation also clarifies attribution. Comparing Transformer-UNI-Argmax with Transformer-PA-Sweep changes both temporal normalization and inference decision logic. The gain should be described as the effect of the combined phase-aligned resampling and sweep-inference configuration. The independent contributions of temporal alignment and sweep inference require further controlled evaluation.",
    )
    insert_heading(doc, anchor6, "5.2. Explainable Scoring and Expert Agreement", 2)
    insert_body_paragraph(
        doc,
        anchor6,
        "The rule-based scoring module provides interpretable feedback aligned with expert-relevant action features such as kicking height, knee extension, recovery, stability, speed, and smoothness. However, all metrics are derived from two-dimensional monocular pose trajectories. The kinematic vigor index should not be interpreted as true biomechanical power.",
    )
    insert_body_paragraph(
        doc,
        anchor6,
        "The expert-annotated subset showed encouraging preliminary agreement, including high inter-expert score reliability and moderate system-expert score association. Nevertheless, the subset contains only 30 videos and cannot currently be verified as independent at the participant level. These results support feasibility but not expert-equivalent scoring.",
    )
    insert_heading(doc, anchor6, "5.3. Limitations and Future Work", 2)
    insert_body_paragraph(
        doc,
        anchor6,
        "The most important limitation is the absence of participant identifiers in the current data files. A leakage-free subject-independent split cannot be created until a participant metadata table is recovered. Future experiments should therefore repeat the full matrix with participant-level train, validation, and test partitions.",
    )
    insert_body_paragraph(
        doc,
        anchor6,
        "Additional limitations include monocular 2D sensing, sensitivity to occlusion and camera viewpoint, possible multi-person interference, classroom-specific recording conditions, and the absence of true smartphone on-device runtime testing. The study evaluates recognition and scoring feasibility; it does not prove improvement in student learning outcomes.",
    )

    delete_section(doc, "6. Conclusion", "Author Contributions:")
    anchor_back = find_para(doc, "Author Contributions:")
    insert_heading(doc, anchor_back, "6. Conclusion", 1)
    add_markdown_section(doc, anchor_back, REV / "conclusion_revised.md", skip_title=True)

    set_para_text(
        doc,
        "Informed Consent Statement:",
        "Informed Consent Statement: Verbal informed consent for participation was obtained from all participants before video recording because the study was conducted as a non-invasive classroom movement-assessment activity and collected derived analytical data for research use. Participants were informed of the study purpose, recording procedure, anonymization measures, voluntary participation, and their right to withdraw without consequences.",
    )
    set_para_text(
        doc,
        "Data Availability Statement:",
        "Data Availability Statement: De-identified derived materials supporting the findings of this study, including pose-keypoint sequences, code, experiment manifests, and generated result files, are included in the submission/reproducibility package or are available from the corresponding author upon reasonable request. Raw videos are not publicly available due to privacy and ethical restrictions related to human participant recordings.",
    )

    # Add missing baseline references if they are not already present.
    ref_text = "\n".join(p.text for p in doc.paragraphs)
    refs_anchor = None
    for p in reversed(doc.paragraphs):
        if p.text.strip().startswith("19. "):
            refs_anchor = p
            break
    last_ref = refs_anchor
    if refs_anchor is not None and "Long Short-Term Memory" not in ref_text:
        last_ref = insert_paragraph_after(
            doc,
            refs_anchor,
            "20. Hochreiter, S.; Schmidhuber, J. Long short-term memory. Neural Comput. 1997, 9, 1735-1780. https://doi.org/10.1162/neco.1997.9.8.1735"
            ,
            style(doc, "MDPI_8.1_references", "Normal"),
        )
    ref_text = "\n".join(p.text for p in doc.paragraphs)
    if refs_anchor is not None and "Semi-Supervised Classification with Graph Convolutional Networks" not in ref_text:
        insert_paragraph_after(
            doc,
            last_ref,
            "21. Kipf, T.N.; Welling, M. Semi-supervised classification with graph convolutional networks. In Proceedings of the International Conference on Learning Representations (ICLR), Toulon, France, 24-26 April 2017.",
            style(doc, "MDPI_8.1_references", "Normal"),
        )

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
